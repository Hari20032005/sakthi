import docx

doc = docx.Document('review 2.docx')
for i, p in enumerate(doc.paragraphs):
    has_drawing = any('drawing' in r._r.xml for r in p.runs)
    if has_drawing:
        print(f"P[{i}] has drawing! Text: '{p.text}'")

doc1 = docx.Document('revive1.docx')
print("\n--- revive1.docx drawings ---")
for i, p in enumerate(doc1.paragraphs):
    has_drawing = any('drawing' in r._r.xml for r in p.runs)
    if has_drawing:
        print(f"P[{i}] has drawing! Text: '{p.text}'")
