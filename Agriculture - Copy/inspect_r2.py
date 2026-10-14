import docx

doc = docx.Document('review 2.docx')
print(f"Total paragraphs in review 2.docx: {len(doc.paragraphs)}")
for i, p in enumerate(doc.paragraphs):
    has_drawing = any('drawing' in r._r.xml for r in p.runs)
    t = p.text.strip()
    draw_str = " [HAS DRAWING]" if has_drawing else ""
    print(f"P[{i}]{draw_str}: {t}")

print(f"\nTotal tables: {len(doc.tables)}")
for idx, tbl in enumerate(doc.tables):
    print(f"Table {idx}: rows={len(tbl.rows)}, cols={len(tbl.columns)}")
    for r_idx, r in enumerate(tbl.rows):
        cells = [c.text.strip().replace('\n', ' ') for c in r.cells]
        print(f"  R{r_idx}: {cells[:4]}")
